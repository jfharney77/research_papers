from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .config import DEFAULT_TEMPLATE, DOCUMENTS_ROOT, LATEX_ROOT, REPO_ROOT, SCRIPT_ROOT
from .models import DocumentManifest, FigureEntry, SectionEntry
from .refextract import extract_references
from .sandbox import run_sandboxed
from .templates import is_buildable, is_valid_template, valid_template_ids
from .utils import slugify, snake_case, text_to_latex


@dataclass
class ConversionOptions:
    template: str
    build_pdf: bool = True
    overwrite: bool = False


class ConversionError(RuntimeError):
    pass


def convert_document(docx_path: Path, options: ConversionOptions | None = None) -> DocumentManifest:
    docx_path = docx_path.expanduser().resolve()
    if not docx_path.exists():
        raise FileNotFoundError(docx_path)

    options = options or ConversionOptions(template=DEFAULT_TEMPLATE, build_pdf=True)
    if not is_valid_template(options.template):
        raise ConversionError(
            f"Unknown template '{options.template}'. Valid templates: "
            f"{', '.join(valid_template_ids()) or '(none found)'}"
        )
    if options.build_pdf and not is_buildable(options.template):
        raise ConversionError(
            f"Template '{options.template}' has no build script "
            f"(expected {SCRIPT_ROOT / options.template / 'build.sh'}). "
            "Re-run with build disabled, or add the build script."
        )
    template_dir = (LATEX_ROOT / options.template).resolve()

    DOCUMENTS_ROOT.mkdir(parents=True, exist_ok=True)

    title = docx_path.stem.replace("_", " ").title()
    document_id = snake_case(docx_path.stem)
    workspace = DOCUMENTS_ROOT / document_id
    if workspace.exists():
        if options.overwrite:
            shutil.rmtree(workspace)
        else:
            raise ConversionError(
                f"Workspace {workspace} already exists. Remove or rename before converting."
            )

    workspace.mkdir(parents=True)
    stored_docx = workspace / docx_path.name
    if docx_path != stored_docx:
        stored_docx.write_bytes(docx_path.read_bytes())

    template_workspace = workspace / options.template
    _copy_template(template_dir, template_workspace)

    manifest = _build_from_docx(
        docx_path=Document(stored_docx),
        source_path=stored_docx,
        workspace=workspace,
        template_name=options.template,
        template_workspace=template_workspace,
    )

    manifest.source_docx = stored_docx.name

    manifest_path = workspace / "manifest.json"
    manifest_dict = manifest.model_dump()
    manifest_path.write_text(json.dumps(manifest_dict, indent=2, default=str))

    if options.build_pdf:
        _run_build(template_name=options.template, template_workspace=template_workspace, manifest_path=manifest_path)

    return manifest


def rebuild_pdf(document_id: str) -> DocumentManifest:
    workspace = DOCUMENTS_ROOT / document_id
    manifest_path = workspace / "manifest.json"
    if not manifest_path.exists():
        raise ConversionError(f"Document '{document_id}' not found")

    manifest_data = json.loads(manifest_path.read_text())
    template_name = manifest_data.get("template")
    template_workspace = workspace / template_name
    if not template_workspace.exists():
        raise ConversionError(
            f"Template workspace missing for '{document_id}'. Expected {template_workspace}"
        )

    _run_build(template_name=template_name, template_workspace=template_workspace, manifest_path=manifest_path)
    refreshed = json.loads(manifest_path.read_text())
    return DocumentManifest(**refreshed)


def _copy_template(template_dir: Path, destination: Path) -> None:
    import shutil

    shutil.copytree(template_dir, destination)


def _build_from_docx(*, docx_path: Document, source_path: Path, workspace: Path, template_name: str, template_workspace: Path) -> DocumentManifest:
    sections_dir = template_workspace / "sections"
    sections_dir.mkdir(exist_ok=True)
    references_dir = template_workspace / "references"
    references_dir.mkdir(exist_ok=True)
    assets_dir = template_workspace / "assets"
    assets_dir.mkdir(exist_ok=True)

    manifest = DocumentManifest(
        document_id=workspace.name,
        title=workspace.name.replace("_", " ").title(),
        template=template_name,
        source_docx="",
    )

    sections: list[SectionEntry] = []
    figures: list[FigureEntry] = []

    current_section_slug: str | None = None
    current_lines: list[str] = []
    order = 1
    figure_counter = 1

    def ensure_section(default_title: str = "Section") -> None:
        nonlocal current_section_slug, order
        if current_section_slug is None:
            section_slug = slugify(f"{default_title}-{order}")
            latex_path = f"sections/{section_slug}.tex"
            sections.append(SectionEntry(
                section_id=section_slug,
                title=f"{default_title} {order}",
                level=1,
                slug=section_slug,
                latex_path=latex_path,
                order=order,
            ))
            current_section_slug = section_slug
            order += 1

    paragraphs: Iterable[Paragraph] = docx_path.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        style_name = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()

        if style_name.startswith("Heading"):
            # flush previous
            if current_section_slug is not None:
                _write_section_file(
                    sections_dir,
                    current_section_slug,
                    current_lines,
                )
                current_lines = []

            level = _heading_level(style_name)
            section_slug = slugify(text or f"section-{order}")
            latex_path = f"sections/{section_slug}.tex"
            sections.append(SectionEntry(
                section_id=section_slug,
                title=text or f"Section {order}",
                level=level,
                slug=section_slug,
                latex_path=latex_path,
                order=order,
            ))
            current_section_slug = section_slug
            order += 1
            if text:
                current_lines.append(_latex_heading(level, text))
            continue

        if _paragraph_contains_figure(paragraph):
            fig_path = _extract_figure(paragraph, assets_dir, figure_counter)
            caption = _next_paragraph_caption(docx_path, idx + 1)
            latex_snippet = _latex_figure_block(fig_path, caption)
            ensure_section()
            current_lines.append(latex_snippet)
            figures.append(FigureEntry(
                figure_id=f"figure-{figure_counter}",
                caption=caption,
                relative_path=str(Path("assets") / Path(fig_path).name),
                section_slug=current_section_slug,
            ))
            figure_counter += 1
            continue

        if text:
            ensure_section()
            current_lines.append(text_to_latex(text) + "\n\n")

    if current_section_slug is not None and current_lines:
        _write_section_file(sections_dir, current_section_slug, current_lines)

    manifest.sections = sections
    manifest.figures = figures

    # Extract the bibliography from the source document and emit a populated
    # references.bib. When extraction is partial or impossible, the warning is
    # surfaced in the manifest instead of silently writing a comment stub.
    #
    # The templates all do \bibliography{references}, which bibtex resolves to
    # references.bib next to main.tex — so the extracted entries replace the
    # placeholder bib in the template root. A copy is also kept under
    # references/ for reference.
    bib_content, references_warning = extract_references(source_path)
    (template_workspace / "references.bib").write_text(bib_content)
    (references_dir / "references.bib").write_text(bib_content)
    manifest.references_warning = references_warning

    _rewrite_main(template_workspace, manifest.sections)

    return manifest


def _heading_level(style_name: str) -> int:
    parts = style_name.split()
    for part in parts:
        if part.isdigit():
            return int(part)
        if part.endswith("Heading") and part[:-7].isdigit():
            return int(part[:-7])
    # default when style names like "Heading 1"
    try:
        return int(parts[-1])
    except (ValueError, IndexError):
        return 1


def _latex_heading(level: int, text: str) -> str:
    latex_text = text_to_latex(text)
    if level <= 1:
        return f"\\section{{{latex_text}}}\n\n"
    if level == 2:
        return f"\\subsection{{{latex_text}}}\n\n"
    return f"\\subsubsection{{{latex_text}}}\n\n"


def _write_section_file(sections_dir: Path, slug: str, lines: list[str]) -> None:
    target = sections_dir / f"{slug}.tex"
    target.write_text("".join(lines))


_NS = {
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def _paragraph_contains_figure(paragraph: Paragraph) -> bool:
    element = etree.fromstring(paragraph._p.xml)
    return bool(element.xpath(".//pic:pic", namespaces=_NS))


def _extract_figure(paragraph: Paragraph, assets_dir: Path, counter: int) -> Path:
    element = etree.fromstring(paragraph._p.xml)
    pics = element.xpath(".//pic:pic", namespaces=_NS)
    if not pics:
        raise ConversionError("Figure expected but not found in paragraph XML.")

    blip = pics[0].xpath(".//a:blip", namespaces=_NS)[0]
    embed_id = blip.get(qn("r:embed"))
    document_part = paragraph.part
    image_part = document_part.related_parts[embed_id]
    ext = Path(image_part.partname).suffix or ".png"
    output = assets_dir / f"figure_{counter:02d}{ext}"
    output.write_bytes(image_part.blob)
    return output


def _next_paragraph_caption(doc: Document, start_index: int) -> str | None:
    paragraphs = doc.paragraphs
    if start_index >= len(paragraphs):
        return None
    next_para = paragraphs[start_index]
    if next_para.style and "Caption" in next_para.style.name:
        return next_para.text.strip() or None
    return None


def _latex_figure_block(image_path: Path, caption: str | None) -> str:
    rel_path = Path("assets") / image_path.name
    latex = ["\\begin{figure}[t]\n", "  \\centering\n"]
    latex.append(f"  \\includegraphics[width=0.85\\linewidth]{{{rel_path.as_posix()}}}\n")
    if caption:
        latex.append(f"  \\caption{{{text_to_latex(caption)}}}\n")
    latex.append("\\end{figure}\n\n")
    return "".join(latex)


def _rewrite_main(template_workspace: Path, sections: list[SectionEntry]) -> None:
    main_tex = template_workspace / "main.tex"
    if not main_tex.exists():
        return

    content = main_tex.read_text()
    if "\\begin{document}" not in content:
        return

    before, after = content.split("\\begin{document}", 1)
    if "\\end{document}" in after:
        body, tail = after.split("\\end{document}", 1)
    else:
        body = after
        tail = ""

    section_inputs = ["% Auto-generated sections\n"]
    for entry in sorted(sections, key=lambda s: s.order):
        section_inputs.append(f"\\input{{sections/{entry.slug}}}\n")

    rebuilt = (
        before
        + "\\begin{document}\n"
        + body
        + "\n"
        + "".join(section_inputs)
        + "% ---- End auto-generated sections ----\n"
        + "\\end{document}\n"
        + tail
    )
    main_tex.write_text(rebuilt)


def _run_build(*, template_name: str, template_workspace: Path, manifest_path: Path) -> None:
    script = SCRIPT_ROOT / template_name / "build.sh"
    log_file = template_workspace / "build.log"
    manifest_data = json.loads(manifest_path.read_text())

    # Compile under the sandbox: timeout, resource limits, no shell-escape, and
    # TeX file IO restricted to the workspace. For docker, the build script and
    # workspace are both mounted at /work.
    result = run_sandboxed(
        ["bash", str(script), str(template_workspace)],
        cwd=REPO_ROOT,
        workspace=template_workspace,
        docker_cmd=["bash", f"/work/{script.name}", "/work"],
    )
    log_file.write_text(result.stdout + "\n" + result.stderr)
    if result.returncode == 0:
        manifest_data["build"] = {
            "status": "succeeded",
            "last_run": datetime.utcnow().isoformat(),
            "pdf_path": f"{template_name}/main.pdf",
            "log_path": str(log_file.relative_to(template_workspace.parent)),
        }
    else:
        message = (
            f"LaTeX build timed out after the configured limit."
            if result.timed_out
            else "LaTeX build failed. See log for details."
        )
        manifest_data["build"] = {
            "status": "failed",
            "last_run": datetime.utcnow().isoformat(),
            "pdf_path": None,
            "log_path": str(log_file.relative_to(template_workspace.parent)),
            "message": message,
        }
    manifest_path.write_text(json.dumps(manifest_data, indent=2))
