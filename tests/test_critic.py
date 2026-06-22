"""Tests for the Critic: heuristics, text prep, providers, pipeline, and API."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from critic.heuristics import compute_heuristics
from critic.pipeline import critique_adhoc
from critic.providers import available_providers, get_provider
from critic.providers.stub import StubProvider
from critic.textprep import sections_from_docx, strip_latex
from docserver.main import app

client = TestClient(app)

AI_SOUNDING = (
    "Moreover, it is important to note that this robust and comprehensive framework "
    "plays a crucial role. Furthermore, it leverages a seamless, scalable, and "
    "holistic approach. Additionally, this underscores the importance of synergy. "
    "Notably, it delves into a rich tapestry of intricate and pivotal concepts."
)

HUMAN_SOUNDING = (
    "We ran the model on 4 GPUs. It crashed twice. The third run finished in nine "
    "hours and produced a 2.3-point gain over the baseline, which surprised us. "
    "Why? The attention maps suggest the encoder simply memorized rare tokens. "
    "We have not fully explained it. A small ablation, removing layer 7, erased "
    "most of the gain."
)


# --- heuristics --------------------------------------------------------------

def test_ai_text_scores_higher_than_human():
    ai = compute_heuristics(AI_SOUNDING).ai_score
    human = compute_heuristics(HUMAN_SOUNDING).ai_score
    assert ai > human
    assert ai >= 50


def test_signals_flag_filler():
    signals = compute_heuristics(AI_SOUNDING).signals
    assert any("filler" in s for s in signals)


def test_short_text_is_not_assessed():
    report = compute_heuristics("Too short.")
    assert report.ai_score == 0


# --- text prep ---------------------------------------------------------------

def test_strip_latex_removes_commands_and_math():
    out = strip_latex(r"\section{Intro} Text with \cite{foo} and $x^2$ math. % comment")
    assert "\\cite" not in out and "$" not in out
    assert "Intro" in out and "Text with" in out
    assert "comment" not in out


# --- providers ---------------------------------------------------------------

def test_stub_returns_exactly_three_criticisms():
    result = StubProvider().analyze_section(rubric="r", prose=AI_SOUNDING, signals=["x"])
    assert len(result.criticisms) == 3
    assert 0 <= result.ai_score_model <= 100
    assert result.deai_tips


def test_default_provider_is_stub(monkeypatch):
    monkeypatch.delenv("CRITIC_PROVIDER", raising=False)
    assert get_provider().name == "stub"


def test_provider_listing():
    listing = available_providers()
    ids = {p.id for p in listing.providers}
    assert {"stub", "claude", "ollama", "cerebras"} <= ids
    assert any(p.id == "stub" and p.available for p in listing.providers)


# --- pipeline (ad-hoc) -------------------------------------------------------

def _make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(AI_SOUNDING)
    doc.add_heading("Method", level=1)
    doc.add_paragraph(HUMAN_SOUNDING)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_sections_from_docx_splits_on_headings():
    sections = sections_from_docx(_make_docx())
    titles = [s.title for s in sections]
    assert "Introduction" in titles and "Method" in titles


def test_critique_adhoc_end_to_end():
    result = critique_adhoc("paper.docx", _make_docx())
    assert result.document_id is None
    assert result.provider == "stub"
    assert len(result.sections) == 2
    for s in result.sections:
        assert len(s.criticisms) == 3
        assert 0 <= s.ai_score <= 100
        assert set(s.ai_score_breakdown) == {"heuristic", "model"}


# --- API ---------------------------------------------------------------------

def test_api_providers_endpoint():
    res = client.get("/critic/providers")
    assert res.status_code == 200
    assert res.json()["active"] == "stub"


def test_api_critique_missing_document():
    assert client.post("/documents/nope/critique").status_code == 404
    assert client.get("/documents/nope/critique").status_code == 404


def test_api_adhoc_rejects_non_pdf_docx():
    files = {"file": ("notes.txt", io.BytesIO(b"hi"), "text/plain")}
    assert client.post("/critic/adhoc", files=files).status_code == 400


def test_api_adhoc_docx():
    files = {"file": ("paper.docx", io.BytesIO(_make_docx()), "application/octet-stream")}
    res = client.post("/critic/adhoc", files=files)
    assert res.status_code == 200
    body = res.json()
    assert len(body["sections"]) == 2
    assert body["provider"] == "stub"
