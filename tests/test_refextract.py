"""Tests for bibliography extraction from .docx (docbuilder.refextract)."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from docbuilder.refextract import extract_references


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_BIB_NS = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"


def _write_docx(tmp_path, build):
    doc = Document()
    build(doc)
    path = tmp_path / "paper.docx"
    doc.save(str(path))
    return path


def _inject_sources_xml(docx_path, sources_xml: str):
    """Add a customXml part containing a <b:Sources> root to an existing .docx."""
    data = docx_path.read_bytes()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buffer, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("customXml/item1.xml", sources_xml)
    docx_path.write_bytes(buffer.getvalue())
    return docx_path


# ---------------------------------------------------------------------------
# Step 1 — Word Sources XML
# ---------------------------------------------------------------------------

def test_sources_xml_produces_valid_entries(tmp_path):
    sources_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<b:Sources xmlns:b="{_BIB_NS}">
  <b:Source>
    <b:Tag>Smith2020</b:Tag>
    <b:SourceType>JournalArticle</b:SourceType>
    <b:Author>
      <b:Author>
        <b:NameList>
          <b:Person><b:Last>Smith</b:Last><b:First>John</b:First></b:Person>
          <b:Person><b:Last>Doe</b:Last><b:First>Jane</b:First></b:Person>
        </b:NameList>
      </b:Author>
    </b:Author>
    <b:Title>A Great Paper</b:Title>
    <b:JournalName>Journal of Stuff</b:JournalName>
    <b:Year>2020</b:Year>
    <b:Volume>5</b:Volume>
    <b:Pages>1-10</b:Pages>
    <b:DOI>10.1234/abc</b:DOI>
  </b:Source>
  <b:Source>
    <b:Tag>Lee2019</b:Tag>
    <b:SourceType>ConferenceProceedings</b:SourceType>
    <b:Author>
      <b:Author>
        <b:NameList>
          <b:Person><b:Last>Lee</b:Last><b:First>Kim</b:First></b:Person>
        </b:NameList>
      </b:Author>
    </b:Author>
    <b:Title>Conference Work</b:Title>
    <b:ConferenceName>Big Conf</b:ConferenceName>
    <b:Year>2019</b:Year>
  </b:Source>
</b:Sources>
"""
    path = _write_docx(tmp_path, lambda d: d.add_paragraph("body"))
    _inject_sources_xml(path, sources_xml)

    bib, warning = extract_references(path)

    assert warning is None
    assert "@article{Smith2020," in bib
    assert "author = {Smith, John and Doe, Jane}" in bib
    assert "title = {A Great Paper}" in bib
    assert "journal = {Journal of Stuff}" in bib
    assert "doi = {10.1234/abc}" in bib
    assert "@inproceedings{Lee2019," in bib
    assert "booktitle = {Big Conf}" in bib


# ---------------------------------------------------------------------------
# Step 2 — heuristic paragraph extraction
# ---------------------------------------------------------------------------

def test_heuristic_parses_numbered_references(tmp_path):
    def build(doc):
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Some intro text.")
        doc.add_heading("References", level=1)
        doc.add_paragraph('[1] J. Smith, "A Great Paper," Journal of Stuff, 2020.')
        doc.add_paragraph('[2] K. Lee, "Conference Work," Big Conf, 2019. 10.1234/abc')

    path = _write_docx(tmp_path, build)
    bib, warning = extract_references(path)

    # Both entries parsed cleanly → no warning.
    assert warning is None
    assert "@misc{ref1," in bib
    assert "title = {A Great Paper}" in bib
    assert "year = {2020}" in bib
    assert "doi = {10.1234/abc}" in bib


def test_heuristic_unparseable_lines_become_stubs(tmp_path):
    def build(doc):
        doc.add_heading("References", level=1)
        doc.add_paragraph("This is just an unstructured note without a quoted title.")
        doc.add_paragraph("Another freeform reference line.")

    path = _write_docx(tmp_path, build)
    bib, warning = extract_references(path)

    assert warning is not None
    assert "could not be parsed" in warning
    assert bib.count("@misc{") == 2
    assert "note = {" in bib


def test_references_section_stops_at_next_heading(tmp_path):
    def build(doc):
        doc.add_heading("References", level=1)
        doc.add_paragraph('[1] A. Author, "Title," Venue, 2021.')
        doc.add_heading("Appendix", level=1)
        doc.add_paragraph("This should not be treated as a reference.")

    path = _write_docx(tmp_path, build)
    bib, _ = extract_references(path)

    assert bib.count("@misc{") == 1
    assert "Appendix" not in bib
    assert "should not be treated" not in bib


# ---------------------------------------------------------------------------
# Empty case
# ---------------------------------------------------------------------------

def test_no_references_sets_warning_and_no_stub(tmp_path):
    def build(doc):
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Body text only, no references at all.")

    path = _write_docx(tmp_path, build)
    bib, warning = extract_references(path)

    assert warning == "No references could be extracted; references.bib is empty."
    # The old "% TODO" stub must not be emitted.
    assert "% TODO" not in bib
    assert "@" not in bib
